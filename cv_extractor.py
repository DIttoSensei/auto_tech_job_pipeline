from docx import Document

def extract_cv(filepath):
    doc = Document(filepath)
    cv_text = []

    # Extract paragraphs
    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            cv_text.append(paragraph.text.strip())

    # Extract tables (skills section)
    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                if cell.text.strip():
                    row_text.append(cell.text.strip())
            if row_text:
                cv_text.append(" | ".join(row_text))

    full_text = "\n".join(cv_text)

    # Structure into dictionary
    cv_data = {
        "name": "Richard Oluwatomiwa Andrew",
        "email": "richardtommyandrew@gmail.com",
        "phone": "+234 902 788 3012",
        "location": "Nigeria (Open to Remote)",
        "github": "github.com/DIttoSensei",
        "linkedin": "linkedin.com/in/richard-andrew-3315a7260",
        "portfolio": "my-portfolio-rouge-alpha-26.vercel.app",
        "full_cv_text": full_text
    }

    return cv_data

cv = extract_cv("cv.docx")

for key, value in cv.items():
    if key != "full_cv_text":
        print(f"{key}: {value}")

print("\n--- FULL CV TEXT ---")
print(cv["full_cv_text"])